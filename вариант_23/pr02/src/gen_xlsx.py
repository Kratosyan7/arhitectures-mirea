"""Excel для ПР2 (вариант 23): таблицы мультиплексора 16→1, T-триггера и счётчика (строка 24)."""
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import counter_model as cm  # noqa: E402

OUT = os.path.join(HERE, "..", "Практика_2_вариант_23.xlsx")
HEAD = PatternFill("solid", fgColor="DDEBF7")
HL = PatternFill("solid", fgColor="E2EFDA")       # активная строка/значение
BORD = Border(*[Side(style="thin")] * 4)
CEN = Alignment(horizontal="center", vertical="center")


def cell(ws, coord, val, fill=None, bold=False):
    c = ws[coord]
    c.value = val
    if fill:
        c.fill = fill
    if bold:
        c.font = Font(bold=True)
    c.border = BORD
    c.alignment = CEN
    return c


wb = Workbook()

# ---- лист 1: мультиплексор 16→1 ----
ws = wb.active
ws.title = "Мультиплексор 16-1"
ws["A1"] = "Мультиплексор с 4 адресными входами (16→1): адрес A3A2A1A0 выбирает вход D0..D15 на выход Y"
ws["A1"].font = Font(bold=True)
hdr = ["A3", "A2", "A1", "A0", "Выход Y"]
for j, h in enumerate(hdr):
    cell(ws, f"{chr(65+j)}3", h, HEAD, True)
for a in range(16):
    r = a + 4
    bits = [(a >> (3 - k)) & 1 for k in range(4)]
    for k in range(4):
        cell(ws, f"{chr(65+k)}{r}", bits[k])
    cell(ws, f"E{r}", f"Y = D{a}", HL)
ws.column_dimensions["E"].width = 14

# ---- лист 2: T-триггер ----
tt = wb.create_sheet("T-триггер")
tt["A1"] = "T-триггер: характеристическая таблица Q(n+1) = Q ⊕ T"
tt["A1"].font = Font(bold=True)
for j, h in enumerate(["T", "Qn", "Q(n+1)", "Режим"]):
    cell(tt, f"{chr(65+j)}3", h, HEAD, True)
rows = [(0, 0, 0, "хранение"), (0, 1, 1, "хранение"), (1, 0, 1, "переключение"), (1, 1, 0, "переключение")]
for i, (T, q, q1, mode) in enumerate(rows):
    r = i + 4
    cell(tt, f"A{r}", T); cell(tt, f"B{r}", q); cell(tt, f"C{r}", q1, HL); cell(tt, f"D{r}", mode)
tt["A9"] = "Таблица возбуждения (из Qn в Q(n+1)):"
tt["A9"].font = Font(bold=True)
for j, h in enumerate(["Qn", "Q(n+1)", "T"]):
    cell(tt, f"{chr(65+j)}10", h, HEAD, True)
exc = [(0, 0, 0), (0, 1, 1), (1, 0, 1), (1, 1, 0)]
for i, (q, q1, T) in enumerate(exc):
    r = i + 11
    cell(tt, f"A{r}", q); cell(tt, f"B{r}", q1); cell(tt, f"C{r}", T)
tt.column_dimensions["D"].width = 16

# ---- лист 3: счётчик (строка 24) ----
ct = wb.create_sheet("Счётчик")
ct["A1"] = "Счётчик по модулю 16, строка 24 табл. 2.11 (вариант 23)"
ct["A1"].font = Font(bold=True)
ct["A2"] = "Таблица состояний"; ct["A2"].font = Font(bold=True)
for j, h in enumerate(["Позиция", "Значение", "Код Q3Q2Q1Q0"]):
    cell(ct, f"{chr(65+j)}3", h, HEAD, True)
for p, v in enumerate(cm.VAL):
    r = p + 4
    cell(ct, f"A{r}", p); cell(ct, f"B{r}", v); cell(ct, f"C{r}", f"{v:04b}")
# таблица переходов справа
ct["E2"] = "Таблица переходов"; ct["E2"].font = Font(bold=True)
for j, h in enumerate(["Текущее", "Условие", "Следующее"]):
    cell(ct, f"{chr(69+j)}3", h, HEAD, True)
for i, (s, cond, n) in enumerate(cm.transitions()):
    r = i + 4
    cell(ct, f"E{r}", f"{s} ({s:04b})"); cell(ct, f"F{r}", cond or "—"); cell(ct, f"G{r}", f"{n} ({n:04b})")
# функции возбуждения
base = 4 + len(cm.transitions()) + 1
ct[f"E{base}"] = "Функции возбуждения newQ (метод пособия):"; ct[f"E{base}"].font = Font(bold=True)
for bit in range(3, -1, -1):
    t, r_ = cm.manual_terms(bit)
    ct[f"E{base + (3 - bit) + 1}"] = f"newQ{bit} = " + " ∨ ".join(t + r_)
for col, w in dict(A=9, B=9, C=13, E=12, F=10, G=12).items():
    ct.column_dimensions[col].width = w

wb.save(OUT)
print("saved", os.path.abspath(OUT))
