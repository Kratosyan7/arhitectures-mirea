"""Сборка отчёта prNN/report.md → prNN/Отчёт_ПР_NN.docx (+ PDF через LibreOffice).

Титульный лист — по образцу «Пример отчета.pdf». Разметка report.md — см. tools/REPORT_FORMAT.md.
    python tools/build_report.py pr03 [pr04 ...] [--pdf]
"""
import os
import re
import subprocess
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EMBLEM = os.path.join(ROOT, "tools", "report_assets", "emblem-000.png")
META = dict(
    institute="Институт Информационных технологий",
    department="Кафедра Прикладной Математики (ПМ)",
    discipline="«Архитектура вычислительных машин и систем»",
    group="ИКБО-41-25",
    student="Мирзоян А.О.",
    teacher="",
    year="2026",
)
FONT = "Times New Roman"
MONO = "Courier New"
TEXT_W = Cm(16.5)


# ---------- низкоуровневые помощники ----------
def set_font(run, size=14, bold=None, italic=None, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:cs"), name)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    return run


def para(doc, text="", align=None, size=14, bold=False, indent=True, space_after=0, spacing=1.5, keep=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = spacing
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(0)
    pf.first_line_indent = Cm(1.25) if indent else Cm(0)
    if keep:
        pf.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if align is None else align
    if text:
        add_inline(p, text, size, bold)
    return p


INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def add_inline(p, text, size=14, bold=False):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            set_font(p.add_run(part[2:-2]), size, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            set_font(p.add_run(part[1:-1]), size - 1, bold=bold, name=MONO)
        else:
            set_font(p.add_run(part), size, bold=bold)


def bottom_border(p):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in dict(val="single", sz="8", space="1", color="000000").items():
        b.set(qn(f"w:{k}"), v)
    bdr.append(b)
    pPr.append(bdr)


def page_number_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    set_font(run, 12)


def cell_border(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        for k, v in dict(val="single", sz="4", space="0", color="000000").items():
            el.set(qn(f"w:{k}"), v)
        borders.append(el)
    tblPr.append(borders)


# ---------- титульный лист ----------
def title_page(doc, num):
    def line(text, bold=False, size=14, before=0, after=0, align=WD_ALIGN_PARAGRAPH.CENTER):
        p = para(doc, "", align=align, indent=False, spacing=1.0, size=size)
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        if text:
            set_font(p.add_run(text), size, bold=bold)
        return p

    p = line("")
    p.add_run().add_picture(EMBLEM, width=Cm(2.6))
    line("МИНОБРНАУКИ РОССИИ", before=6, after=8)
    line("Федеральное государственное бюджетное образовательное учреждение")
    line("высшего образования")
    line("«МИРЭА – Российский технологический университет»", bold=True)
    bottom_border(line("РТУ МИРЭА", bold=True, after=4))
    line(META["institute"], before=6)
    line(META["department"], before=12)
    line(f"ОТЧЕТ ПО ПРАКТИЧЕСКОЙ РАБОТЕ № {num}", bold=True, before=96)
    line("по дисциплине", bold=True)
    line(META["discipline"], bold=True)
    t = doc.add_table(rows=2, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    rows = [(f"Выполнил студент группы\n{META['group']}", META["student"]),
            ("Принял преподаватель", META["teacher"] or "____________")]
    for r, (left, right) in enumerate(rows):
        for c, text in enumerate((left, right)):
            cell = t.cell(r, c)
            cell.width = Cm(10.5) if c == 0 else Cm(6)
            cp = cell.paragraphs[0]
            cp.alignment = WD_ALIGN_PARAGRAPH.LEFT if c == 0 else WD_ALIGN_PARAGRAPH.RIGHT
            if r == 1 and c == 0:
                cp.paragraph_format.left_indent = Cm(0.7)
            cp.paragraph_format.space_before = Pt(110 if r == 0 else 36)
            for k, ln in enumerate(text.split("\n")):
                run = cp.add_run(ln)
                set_font(run, 14)
                if k < len(text.split("\n")) - 1:
                    run.add_break()
    line(f"Москва {META['year']}", size=13, before=90)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---------- тело ----------
def picture(doc, base, path, caption):
    full = os.path.join(base, path)
    if not os.path.exists(full):
        para(doc, f"[нет файла рисунка: {path}]", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    else:
        w, h = Image.open(full).size
        width = TEXT_W
        if h / w * 16.5 > 21:          # слишком высокая картинка — ограничиваем высоту
            width = Cm(21 * w / h)
        p = para(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, spacing=1.0, keep=True)
        p.paragraph_format.space_before = Pt(6)
        p.add_run().add_picture(full, width=width)
    para(doc, caption, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, size=12, space_after=8, spacing=1.0)


def table(doc, rows, caption):
    if caption:
        para(doc, caption, align=WD_ALIGN_PARAGRAPH.LEFT, indent=False, size=12, spacing=1.0, keep=True)
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_border(t)
    size = 12 if ncol <= 6 else (10 if ncol <= 12 else 8)
    # Шапка может группировать столбцы: ячейка «<» объединяется с соседней слева
    # (горизонтальный colspan), так в шапке видно, какие столбцы к какому полю относятся.
    for i, r in enumerate(rows):
        j = 0
        while j < ncol:
            text = r[j] if j < len(r) else ""
            span = 1
            while j + span < ncol and j + span < len(r) and r[j + span].strip() == "<":
                span += 1
            if span > 1:
                anchor = t.cell(i, j).merge(t.cell(i, j + span - 1))
            else:
                anchor = t.cell(i, j)
            while len(anchor.paragraphs) > 1:                 # убрать лишние абзацы после merge
                el = anchor.paragraphs[-1]._element
                el.getparent().remove(el)
            cp = anchor.paragraphs[0]
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.first_line_indent = Cm(0)
            cp.paragraph_format.space_after = Pt(0)
            cp.paragraph_format.line_spacing = 1.0
            add_inline(cp, text, size, bold=(i == 0 or span > 1))
            j += span
    para(doc, "", spacing=1.0)


def listing(doc, lines, caption):
    if caption:
        para(doc, caption, align=WD_ALIGN_PARAGRAPH.LEFT, indent=False, size=12, spacing=1.0, keep=True)
    t = doc.add_table(rows=1, cols=1)
    cell_border(t)
    cell = t.cell(0, 0)
    cp = cell.paragraphs[0]
    cp.paragraph_format.first_line_indent = Cm(0)
    cp.paragraph_format.line_spacing = 1.0
    cp.paragraph_format.space_after = Pt(0)
    size = 10 if max((len(s) for s in lines), default=0) <= 90 else 8
    for k, ln in enumerate(lines):
        run = set_font(cp.add_run(ln.replace("\t", "    ")), size, name=MONO)
        if k < len(lines) - 1:
            run.add_break()
    para(doc, "", spacing=1.0)


def parse_table(lines):
    rows = []
    for ln in lines:
        cells = [c.strip() for c in ln.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
            continue
        rows.append(cells)
    return rows


def build(prdir, pdf=False):
    base = os.path.join(ROOT, prdir)
    src = open(os.path.join(base, "report.md"), encoding="utf-8").read()
    m = re.search(r"<!--\s*num:\s*([^;]+);\s*title:\s*(.*?)\s*-->", src)
    num = m.group(1).strip() if m else prdir[2:].lstrip("0")
    src = re.sub(r"<!--.*?-->", "", src, flags=re.S)

    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21)
    sec.left_margin, sec.right_margin = Cm(3), Cm(1.5)
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
    page_number_footer(sec)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = FONT, Pt(14)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_after = Pt(0)
    title_page(doc, num)

    lines = src.splitlines()
    i, pending_caption, buf = 0, None, []

    def flush():
        if buf:
            para(doc, " ".join(s.strip() for s in buf))
            buf.clear()

    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if s.startswith("```"):
            flush()
            j = i + 1
            code = []
            while j < len(lines) and not lines[j].strip().startswith("```"):
                code.append(lines[j])
                j += 1
            listing(doc, code, pending_caption)
            pending_caption, i = None, j + 1
            continue
        if s.startswith("|"):
            flush()
            j = i
            while j < len(lines) and lines[j].strip().startswith("|"):
                j += 1
            table(doc, parse_table(lines[i:j]), pending_caption)
            pending_caption, i = None, j
            continue
        if pending_caption and s:
            para(doc, pending_caption, align=WD_ALIGN_PARAGRAPH.LEFT, indent=False, size=12, spacing=1.0)
            pending_caption = None
        if not s:
            flush()
        elif s.startswith("# "):
            flush()
            h = para(doc, s[2:].strip().upper(), align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, indent=False,
                     space_after=6, keep=True)
            h.paragraph_format.space_before = Pt(12)
        elif s.startswith("## ") or s.startswith("### "):
            flush()
            h = para(doc, s.lstrip("#").strip(), bold=True, space_after=4, keep=True,
                     align=WD_ALIGN_PARAGRAPH.LEFT)
            h.paragraph_format.space_before = Pt(8)
        elif re.match(r"!\[.*\]\(.*\)", s):
            flush()
            mm = re.match(r"!\[(.*)\]\((.*)\)", s)
            picture(doc, base, mm.group(2), mm.group(1))
        elif re.match(r"(Таблица|Листинг)\s+[\dА-Я.]+\s+[–-]", s):
            flush()
            pending_caption = s
        elif s.startswith("$$") and s.endswith("$$"):
            flush()
            para(doc, s.strip("$").strip(), align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
        elif re.match(r"[-*] ", s):
            flush()
            p = para(doc, "– " + s[2:].strip())
        elif re.match(r"\d+[.)] ", s):
            flush()
            para(doc, s)
        else:
            buf.append(s)
        i += 1
    flush()
    out = os.path.join(base, f"Отчёт_ПР_{int(num.split('-')[0]):02d}.docx" if num.split('-')[0].isdigit()
                       else f"Отчёт_{prdir}.docx")
    doc.save(out)
    print("wrote", out)
    if pdf:
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", base, out],
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, check=False)
        print("pdf", out[:-5] + ".pdf")
    return out


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    for d in args:
        build(d, pdf="--pdf" in sys.argv)
