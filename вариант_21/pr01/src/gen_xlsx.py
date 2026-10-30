"""Excel для ПР1 (вариант 21): таблица истинности F=E0E1, СКНФ, МДНФ, карты Карно.
Подсвечены строки F=1 и группы единиц в карте Карно (как в примере «заимствованное»)."""
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "..", "Практика_1_вариант_21.xlsx")

VEC = 0xE0E1
F = [(VEC >> (15 - i)) & 1 for i in range(16)]
GRAY = [0, 1, 3, 2]  # 00,01,11,10

GREEN = PatternFill("solid", fgColor="E2EFDA")   # строки F=1
PEACH = PatternFill("solid", fgColor="FFF2CC")   # столбец F
YEL = PatternFill("solid", fgColor="FFFF00")     # заголовки карт
G1 = PatternFill("solid", fgColor="A9D08E")      # группа ¬b¬d
G2 = PatternFill("solid", fgColor="9DC3E6")      # группа ¬b¬c
G3 = PatternFill("solid", fgColor="FFC000")      # группа abcd
BORD = Border(*[Side(style="thin")] * 4)
CEN = Alignment(horizontal="center", vertical="center")

wb = Workbook()
ws = wb.active
ws.title = "Лист1"


def put(coord, val, fill=None, bold=False, border=True, align=True):
    c = ws[coord]
    c.value = val
    if fill:
        c.fill = fill
    if bold:
        c.font = Font(bold=True)
    if border and not isinstance(val, str):
        c.border = BORD
    if align:
        c.alignment = CEN
    return c


# --- таблица истинности (A..E, строки 1..16) ---
for i in range(16):
    r = i + 1
    bits = [(i >> (3 - k)) & 1 for k in range(4)]
    for k, col in enumerate("ABCD"):
        put(f"{col}{r}", bits[k], GREEN if F[i] else None)
    put(f"E{r}", F[i], PEACH)
ws["A1"].value  # headers go above? keep like образец (no header row, значения с 1)

# --- формулы СКНФ / МДНФ (столбец F/G) ---
NAMES = "abcd"
zeros = [i for i in range(16) if not F[i]]
ones = [i for i in range(16) if F[i]]
def bitlist(i):
    return [(i >> (3 - k)) & 1 for k in range(4)]
sknf = " ∧ ".join("(" + " ∨ ".join(("¬" + v if b else v) for v, b in zip(NAMES, bitlist(i))) + ")" for i in zeros)
put("F1", "СКНФ:", bold=True, border=False)
put("F2", "F = " + sknf, border=False, align=False)
put("F3", "МДНФ:", bold=True, border=False)
put("F4", "F = ¬b∧¬d ∨ ¬b∧¬c ∨ a∧b∧c∧d", border=False, align=False)
put("F6", "СДНФ (для сверки):", bold=True, border=False)
sdnf = " ∨ ".join("(" + " ∧ ".join(("" if b else "¬") + v for v, b in zip(NAMES, bitlist(i))) + ")" for i in ones)
put("F7", "F = " + sdnf, border=False, align=False)

# --- карта Карно: значения F ---
def kmap(top_label_row, title, groups=None):
    r0 = top_label_row
    put(f"O{r0}", title, bold=True, border=False)
    hdr = ["AB/CD", "00", "01", "11", "10"]
    for c, h in zip("OPQRS", hdr):
        put(f"{c}{r0+1}", h, YEL)
    rownames = ["00", "01", "11", "10"]
    for ri, rg in enumerate(GRAY):
        rr = r0 + 2 + ri
        put(f"O{rr}", rownames[ri], YEL)
        for ci, cg in enumerate(GRAY):
            col = "PQRS"[ci]
            v = F[(rg << 2) | cg]
            fill = None
            if groups:
                fill = groups.get((ri, ci))
            put(f"{col}{rr}", v, fill)


kmap(1, "Карта Карно (значения F)")
# МДНФ: ряд ab (ri) 0..3 -> 00,01,11,10; столбец cd (ci) -> 00,01,11,10
#  ¬b¬d: углы (ri0,ci0),(ri0,ci3),(ri3,ci0),(ri3,ci3)
#  ¬b¬c: (ri0,ci1),(ri3,ci1) — незанятые углами клетки (ci0 общие с ¬b¬d)
#  abcd: (ri2,ci2)
groups = {(0, 0): G1, (0, 3): G1, (3, 0): G1, (3, 3): G1,
          (0, 1): G2, (3, 1): G2, (2, 2): G3}
kmap(9, "МДНФ: группы единиц (¬b∧¬d — зел., ¬b∧¬c — син., a∧b∧c∧d — оранж.)", groups)

# --- разложение вектора ---
tet = [(VEC >> (12 - 4 * k)) & 0xF for k in range(4)]
put("A18", f"F16 = E0E1 = " + " ".join(f"{t:04b}" for t in tet) + "₂", bold=True, border=False)
put("A19", "Старший бит вектора соответствует набору 0000.", border=False)

for col, w in dict(A=4, B=4, C=4, D=4, E=5, F=10, O=8, P=5, Q=5, R=5, S=5).items():
    ws.column_dimensions[col].width = w
ws.sheet_view.showGridLines = True
wb.save(OUT)
print("saved", os.path.abspath(OUT))
