"""Excel для ПР3 (вариант 21): управляющее слово ОУ для INPR ← MAR + OUTR."""
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "..", "Практика_3_вариант_21.xlsx")

HEAD = PatternFill("solid", fgColor="DDEBF7")
BORD = Border(*[Side(style="thin")] * 4)
CEN = Alignment(horizontal="center", vertical="center")

wb = Workbook()
ws = wb.active
ws.title = "Управляющее слово"


def put(coord, val, fill=None, bold=False, border=True):
    c = ws[coord]
    c.value = val
    if fill:
        c.fill = fill
    if bold:
        c.font = Font(bold=True)
    if border and fill:
        c.border = BORD
    c.alignment = CEN if border else Alignment(horizontal="left", vertical="center")
    return c


put("A1", "Практика №3. Вариант 21", bold=True, border=False)
put("A2", "INPR ← MAR + OUTR", bold=True, border=False)

# группы полей
ws.merge_cells("A4:C4"); put("A4", "MS(A)", HEAD, True)
ws.merge_cells("D4:F4"); put("D4", "MS(B)", HEAD, True)
ws.merge_cells("G4:L4"); put("G4", "ALU", HEAD, True)
ws.merge_cells("M4:P4"); put("M4", "DC", HEAD, True)
put("Q4", "MS(M)", HEAD, True)
put("R4", "WR", HEAD, True)

labels = ["a2", "a1", "a0", "a2", "a1", "a0", "F", "M", "S2", "S1", "S0", "C",
          "E", "x2", "x1", "x0", "a0", "wr"]
cols = "ABCDEFGHIJKLMNOPQR"
for col, lab in zip(cols, labels):
    put(f"{col}5", lab, HEAD, True)

# MS(A)=011(MAR) MS(B)=110(OutR) ALU=000010(A+B) DC=1 010(InpR) MS(M)=0 WR=0
word = [0, 1, 1,  1, 1, 0,  0, 0, 0, 0, 1, 0,  1, 0, 1, 0,  0, 0]
for col, b in zip(cols, word):
    put(f"{col}6", b, None, False)
    ws[f"{col}6"].border = BORD

bits = "".join(map(str, word))
put("A8", "Слово: 011 | 110 | 000010 | 1010 | 0 | 0 = " + bits + "₂ = 1E0A8₁₆", bold=True, border=False)
put("A9", "MS(A)=011: MAR; MS(B)=110: OutR; ALU=000010: A + B (сумма); "
          "DC=1010: E=1, выход 010 → InpR; MS(M)=0 (на INPR не влияет); WR=0: запись в память не выполняется.",
    border=False)
put("A11", "Проверка (модель Logisim pr03_ou.circ): MAR = 0A2B, OutR = 2C3D → после такта InpR = 3668", border=False)

# второй лист — коды источников/приёмников
ws2 = wb.create_sheet("Коды")
ws2["A1"] = "Код"; ws2["B1"] = "Вход MS A / MS B"; ws2["C1"] = "Выход DC (запись в регистр)"
for c in ("A1", "B1", "C1"):
    ws2[c].fill = HEAD; ws2[c].font = Font(bold=True); ws2[c].border = BORD
rows = [("000", "PC", "—"), ("001", "AC", "OutR"), ("010", "IR", "InpR"), ("011", "MAR", "MDR"),
        ("100", "MDR", "MAR"), ("101", "InpR", "IR"), ("110", "OutR", "AC"), ("111", "Ввод", "PC")]
for i, (k, a, d) in enumerate(rows):
    r = i + 2
    ws2[f"A{r}"] = k; ws2[f"B{r}"] = a; ws2[f"C{r}"] = d
    for col in "ABC":
        ws2[f"{col}{r}"].border = BORD
        ws2[f"{col}{r}"].alignment = CEN
ws2.column_dimensions["B"].width = 18
ws2.column_dimensions["C"].width = 28

for col in "ABCDEFGHIJKLMNOPQR":
    ws.column_dimensions[col].width = 4.5
wb.save(OUT)
print("saved", os.path.abspath(OUT))
